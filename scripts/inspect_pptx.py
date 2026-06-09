#!/usr/bin/env python3
"""Inspect a PPTX for basic executive-deck QA signals."""

from __future__ import annotations

import argparse
from pathlib import Path
from pptx import Presentation


def inspect(path: Path) -> dict:
    prs = Presentation(str(path))
    min_font = None
    low_text_slides = []
    total_text_shapes = 0
    titles = []

    for idx, slide in enumerate(prs.slides, start=1):
        text_shapes = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            content = shape.text_frame.text.strip()
            if content:
                text_shapes.append(content)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size:
                        pt = run.font.size.pt
                        min_font = pt if min_font is None else min(min_font, pt)
        total_text_shapes += len(text_shapes)
        if len(text_shapes) < 3:
            low_text_slides.append({"slide": idx, "text_shapes": len(text_shapes)})
        if text_shapes:
            titles.append({"slide": idx, "first_text": text_shapes[0][:120]})

    return {
        "path": str(path),
        "slides": len(prs.slides),
        "min_font_size": min_font,
        "text_shapes": total_text_shapes,
        "low_text_slides": low_text_slides,
        "titles_preview": titles[:10],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    args = parser.parse_args()
    result = inspect(args.pptx)
    for key, value in result.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()

